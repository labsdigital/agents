import json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Load AST
with open('out/_build/ast.json') as f:
    ast = json.load(f)
with open('out/_build/theme.json') as f:
    theme = json.load(f)

# Create document
doc = Document()

# Set margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Process elements
for elem in ast:
    if elem['type'] == 'chapter':
        p = doc.add_heading(elem['title'], level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif elem['type'] == 'h2':
        text = ''.join([r['t'] for r in elem.get('runs', [])])
        p = doc.add_heading(text, level=2)
    elif elem['type'] == 'h3':
        text = ''.join([r['t'] for r in elem.get('runs', [])])
        p = doc.add_heading(text, level=3)
    elif elem['type'] == 'p':
        text = ''.join([r['t'] for r in elem.get('runs', [])])
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif elem['type'] == 'ul':
        for item in elem.get('items', []):
            text = ''.join([r['t'] for r in item.get('runs', [])])
            doc.add_paragraph(text, style='List Bullet')
    elif elem['type'] == 'ol':
        for item in elem.get('items', []):
            text = ''.join([r['t'] for r in item.get('runs', [])])
            doc.add_paragraph(text, style='List Number')
    elif elem['type'] == 'table':
        # Handle table
        rows_data = elem.get('rows', [])
        if rows_data:
            ncols = len(rows_data[0]) if rows_data[0] else 1
            table = doc.add_table(rows=1, cols=ncols)
            table.style = 'Table Grid'
            # Header
            header = rows_data[0]
            for i, cell in enumerate(header):
                table.rows[0].cells[i].text = cell.get('text', '') if isinstance(cell, dict) else str(cell)
            # Data rows
            for row_data in rows_data[1:]:
                row = table.add_row()
                for i, cell in enumerate(row_data):
                    row.cells[i].text = cell.get('text', '') if isinstance(cell, dict) else str(cell)
    elif elem['type'] == 'code':
        lang = elem.get('lang', '')
        code = elem.get('text', '')
        p = doc.add_paragraph(f'[{lang}]')
        p = doc.add_paragraph(code)
        p.style = doc.styles['No Spacing']
    elif elem['type'] == 'callout':
        text = ''.join([r['t'] for r in elem.get('runs', [])])
        callout_type = elem.get('callout_type', 'note')
        title = elem.get('title', '')
        p = doc.add_paragraph()
        p.add_run(f'[{callout_type.upper()}] {title}').bold = True
        doc.add_paragraph(text)

# Save
output = 'out/Pembelajaran-dan-Asesmen.docx'
doc.save(output)
print(f'DOCX created: {output}')
